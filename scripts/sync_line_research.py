#!/usr/bin/env python3
"""Read LINE research with the SAME service-account token as broker reports.
Dry run is the default. Never sends notifications or calls an AI service.
"""
import calendar
from io import BytesIO
from openpyxl import load_workbook
import hashlib
import json
import os
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from urllib.parse import quote

import requests

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'data/line_research.json'
BROKER_FILE = ROOT / 'data/reports.json'
SHEET_ID = os.getenv('LINE_SHEET_ID', '1ujXlmxd7h5lap-Np6saR9i9Dp1V8Q5orgO9Ko1ccISM').strip()
SHEET_TAB = os.getenv('LINE_SHEET_TAB', 'LINE產業研究資料庫').strip()
COLUMNS = ('date','type','source','industry','company','ticker','period','story',
           'revenue','eps','inventory','utilization','guidance','trend','rating_target',
           'numbers','dividend_current','dividend_next','dividend_policy',
           'source_row','full_text','full_text_2','full_text_3')
BROKER_TYPES = {'券商評等', '個股報告', '產業報告'}


def normalize(value):
    return re.sub(r'[^\w\u4e00-\u9fff]', '', unicodedata.normalize('NFKC', str(value or '')).casefold())


def cutoff_date(today):
    y, m = today.year, today.month - 3
    if m <= 0:
        y -= 1
        m += 12
    return date(y, m, min(today.day, calendar.monthrange(y, m)[1]))


def parse_date(value):
    match = re.search(r'(20\d{2})[-/年](\d{1,2})[-/月](\d{1,2})', str(value or ''))
    if not match:
        return None
    try:
        return date(*map(int, match.groups()))
    except ValueError:
        return None


def load_json(path):
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else {}


def ticker(value):
    match = re.search(r'(?<!\d)\d{4}(?!\d)', str(value or ''))
    return match.group() if match else ''


def broker_candidates(row, brokers):
    if row['type'] not in BROKER_TYPES:
        return []
    stock_code = ticker(row['ticker'])
    company = normalize(row['company'])
    matches = []
    for report in brokers:
        if report.get('date') != row['date']:
            continue
        same_stock = ((stock_code and ticker(report.get('ticker')) == stock_code)
                      or (company and normalize(report.get('name')) == company))
        if same_stock:
            matches.append({'id': report.get('id'), 'broker': report.get('broker'),
                            'title': report.get('title')})
    return matches


def transform(values, brokers, today):
    if not values or not values[0]:
        raise ValueError('Empty sheet; refuse to overwrite existing data')
    cutoff = cutoff_date(today)
    items, seen = [], set()
    skipped_old = skipped_invalid = duplicates = 0
    for sheet_row_number, raw in enumerate(values[1:], start=2):
        cells = (list(raw) + [''] * len(COLUMNS))[:len(COLUMNS)]
        record = dict(zip(COLUMNS, cells))
        day = parse_date(record['date'])
        if day is None:
            skipped_invalid += 1
            continue
        if day < cutoff or day > today:
            skipped_old += 1
            continue
        record['date'] = day.isoformat()
        record['sheet_row_number'] = sheet_row_number
        content = json.dumps({key: record[key] for key in COLUMNS}, ensure_ascii=False, sort_keys=True)
        digest = hashlib.sha256(content.encode('utf-8')).hexdigest()
        if digest in seen:
            duplicates += 1
            continue
        seen.add(digest)
        record['id'] = 'line-' + digest[:20]
        record['possible_broker_duplicates'] = broker_candidates(record, brokers)
        items.append(record)
    items.sort(key=lambda r: (r['date'], r['sheet_row_number']), reverse=True)
    return {'window_start': cutoff.isoformat(), 'window_end': today.isoformat(),
            'count': len(items), 'items': items}, {
            'input_rows': len(values)-1, 'kept': len(items), 'skipped_outside_window': skipped_old,
            'invalid_dates': skipped_invalid, 'identical_rows_removed': duplicates,
            'broker_overlap_candidates': sum(bool(r['possible_broker_duplicates']) for r in items)}


XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
GOOGLE_SHEET_MIME = 'application/vnd.google-apps.spreadsheet'


def drive_get(url, token, params, stage):
    try:
        response = requests.get(url, headers={'Authorization': 'Bearer ' + token},
                                params=params, timeout=120)
    except requests.RequestException:
        raise RuntimeError(stage + ': Google Drive connection failed') from None
    if response.ok:
        return response
    try:
        err = response.json().get('error', {})
    except (ValueError, AttributeError):
        err = {}
    reasons = [e.get('reason') for e in err.get('errors', []) if isinstance(e, dict)]
    reasons += [e.get('reason') for e in err.get('details', []) if isinstance(e, dict)]
    print('DRIVE_API_ERROR ' + json.dumps({
        'stage': stage, 'http_status': response.status_code,
        'status': err.get('status'), 'reasons': reasons
    }), file=sys.stderr)
    # Never log tokens, request headers, workbook contents or arbitrary server messages.
    if 'exportSizeLimitExceeded' in reasons:
        hint = 'Google 試算表匯出超過 10 MB 上限；請提供較小的匯出檔案'
    elif 'SERVICE_DISABLED' in reasons or 'accessNotConfigured' in reasons:
        hint = '請在授權專案啟用 Google Drive API'
    elif response.status_code in (403, 404):
        hint = '此帳號無法執行此檔案操作；請依 stage 區分中繼資料或下載權限，不代表一定未分享'
    elif response.status_code == 401:
        hint = '權杖無效或過期'
    else:
        hint = 'Google Drive request failed'
    raise RuntimeError(stage + ': ' + hint)


def read_xlsx(content):
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise RuntimeError('INVALID_XLSX: 下載內容不是可讀取的 Excel 活頁簿') from None
    try:
        if SHEET_TAB not in workbook.sheetnames:
            raise RuntimeError('SHEET_TAB_NOT_FOUND: 檔案已下載，但找不到指定分頁 ' + SHEET_TAB)
        values = []
        for raw in workbook[SHEET_TAB].iter_rows(max_col=len(COLUMNS), values_only=True):
            row = []
            for value in raw:
                if isinstance(value, (datetime, date)):
                    value = value.isoformat()
                elif value is None:
                    value = ''
                row.append(value)
            values.append(row)
        if not values or not any(v != '' for v in values[0]):
            raise RuntimeError('EMPTY_WORKSHEET: 指定分頁無標題資料，停止匯入')
        print('DRIVE_WORKBOOK_READ_OK: 指定分頁已解析；公式使用檔案儲存的計算結果')
        return values
    finally:
        workbook.close()


def fetch_sheet(token):
    if not re.fullmatch(r'[A-Za-z0-9_-]+', SHEET_ID):
        raise ValueError('LINE_SHEET_ID 必須填檔案 ID，不是完整網址')
    file_id = SHEET_ID
    visited = set()
    # Follow Drive shortcuts only through Google's API and retain target permission checks.
    for _ in range(5):
        if file_id in visited:
            raise RuntimeError('DRIVE_SHORTCUT_LOOP')
        visited.add(file_id)
        base = 'https://www.googleapis.com/drive/v3/files/' + quote(file_id, safe='')
        metadata = drive_get(base, token, {
            'fields': 'id,mimeType,capabilities(canDownload),shortcutDetails(targetId)',
            'supportsAllDrives': 'true'
        }, 'METADATA').json()
        mime = metadata.get('mimeType')
        if mime != 'application/vnd.google-apps.shortcut':
            break
        file_id = metadata.get('shortcutDetails', {}).get('targetId', '')
        if not re.fullmatch(r'[A-Za-z0-9_-]+', file_id):
            raise RuntimeError('INVALID_SHORTCUT_TARGET')
    else:
        raise RuntimeError('DRIVE_SHORTCUT_DEPTH_EXCEEDED')
    print('DRIVE_FILE_ACCESS_OK ' + json.dumps({
        'mimeType': mime, 'canDownload': metadata.get('capabilities', {}).get('canDownload')
    }))
    if metadata.get('capabilities', {}).get('canDownload') is False:
        raise RuntimeError('DOWNLOAD_NOT_ALLOWED: Google Drive 明確禁止此帳號下載；停止')
    if mime == GOOGLE_SHEET_MIME:
        response = drive_get(base + '/export', token, {'mimeType': XLSX_MIME}, 'EXPORT_XLSX')
    elif mime in (XLSX_MIME, 'application/vnd.ms-excel.sheet.macroEnabled.12'):
        response = drive_get(base, token, {'alt': 'media', 'supportsAllDrives': 'true'}, 'DOWNLOAD_XLSX')
    else:
        raise RuntimeError('UNSUPPORTED_FILE_TYPE: ' + str(mime) + '；請提供 .xlsx 或 Google 試算表')
    return read_xlsx(response.content)


def main():
    today = (date.fromisoformat(os.environ['LINE_SYNC_TODAY'])
             if os.getenv('LINE_SYNC_TODAY') else datetime.now(ZoneInfo('Asia/Taipei')).date())
    fixture = os.getenv('LINE_SYNC_FIXTURE')
    if fixture:
        values = json.loads(Path(fixture).read_text(encoding='utf-8'))
    else:
        token = os.getenv('ACCESS_TOKEN', '').strip()
        if not token:
            raise RuntimeError('Missing ACCESS_TOKEN from google-github-actions/auth; check workflow WIF setup')
        values = fetch_sheet(token)
    brokers = load_json(BROKER_FILE).get('items', [])
    data, stats = transform(values, brokers, today)
    print('LINE_SYNC_STATS ' + json.dumps(stats, ensure_ascii=False))
    if os.getenv('LINE_SYNC_DRY_RUN', '1') == '1':
        print('DRY_RUN: no data written, no AI calls, no Telegram')
        return
    # Explicit opt-in required to avoid exposing research text in a public repository.
    if os.getenv('LINE_ALLOW_PUBLIC_WRITE') != 'I_UNDERSTAND':
        raise RuntimeError('Refusing to write research text without explicit LINE_ALLOW_PUBLIC_WRITE opt-in')
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if load_json(OUTPUT) == data:
        print('Unchanged; skip write')
        return
    tmp = OUTPUT.with_suffix('.json.tmp')
    tmp.write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    tmp.replace(OUTPUT)
    print('Updated LINE data locally; no AI calls or notifications')


if __name__ == '__main__':
    main()

